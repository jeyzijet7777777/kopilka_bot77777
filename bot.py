import asyncio
import calendar
import datetime
import io
import logging
import math
import os
import re
import sqlite3
from contextlib import contextmanager
from zoneinfo import ZoneInfo

import matplotlib
matplotlib.use("Agg")
from matplotlib.colors import to_rgba
from matplotlib.figure import Figure
from matplotlib.lines import Line2D
from matplotlib.patches import Patch, Rectangle
from matplotlib.ticker import FuncFormatter
import openpyxl

from telegram import (
    BotCommand,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    InputMediaPhoto,
    ReplyKeyboardMarkup,
    Update,
)
from telegram.error import Forbidden, TelegramError
from telegram.ext import (
    Application,
    CallbackQueryHandler,
    CommandHandler,
    ContextTypes,
    MessageHandler,
    PicklePersistence,
    filters,
)
from telegram.helpers import escape_markdown


# --- Загрузка настроек из .env ---
def load_env(path=".env"):
    if os.path.exists(path):
        with open(path, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if line and not line.startswith("#") and "=" in line:
                    k, v = line.split("=", 1)
                    os.environ.setdefault(k.strip(), v.strip().strip("'").strip('"'))


load_env()

BOT_TOKEN = os.environ.get("BOT_TOKEN", "")
TZ_NAME = os.environ.get("TZ_NAME", "Europe/Moscow")
DB_PATH = os.environ.get("DB_PATH", "budget.db")
STATE_PATH = os.environ.get("STATE_PATH", "bot_state.pkl")
TEMPLATE_PATH = "template.xlsx"

TZ = ZoneInfo(TZ_NAME)
TZ_LABEL = "МСК" if TZ_NAME == "Europe/Moscow" else TZ_NAME
MAX_AMOUNT = 10_000_000

logging.basicConfig(
    format="%(asctime)s %(levelname)s %(name)s: %(message)s",
    level=logging.INFO,
)
# httpx пишет в лог полные URL запросов, а в них лежит токен бота
logging.getLogger("httpx").setLevel(logging.WARNING)
logging.getLogger("httpcore").setLevel(logging.WARNING)
logger = logging.getLogger("budget-bot")

# --- Цвета графиков ---
BLUE = "#2a78d6"
ORANGE = "#eb6834"
GREEN = "#1baf7a"
GREEN_D = "#16a34a"
RED = "#e34948"
GRAY = "#898781"
INK = "#1f2937"

MONTHS_RU = [
    "Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
    "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь",
]

INCOME_MODES = {
    "daily": "📅 Каждый день (например, 1000₽/день)",
    "monthly": "📅 1 раз в месяц (зарплата)",
    "twice_monthly": "📅 2 раза в месяц (аванс + зп)",
    "weekly": "📅 Раз в неделю",
}
INCOME_MODE_SHORT = {
    "daily": "в день",
    "monthly": "в месяц",
    "twice_monthly": "2 раза в месяц",
    "weekly": "в неделю",
}
INCOME_UNIT_PROMPT = {
    "daily": "в день",
    "monthly": "в месяц",
    "twice_monthly": "за один раз (выплат две в месяц)",
    "weekly": "в неделю",
}


# --- База данных ---
@contextmanager
def db():
    """Соединение с коммитом и гарантированным закрытием."""
    conn = sqlite3.connect(DB_PATH, timeout=10)
    conn.row_factory = sqlite3.Row
    try:
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def init_db():
    boot = sqlite3.connect(DB_PATH)
    boot.execute("PRAGMA journal_mode=WAL")
    boot.close()

    with db() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS settings (
                user_id INTEGER,
                key TEXT,
                value TEXT,
                PRIMARY KEY (user_id, key)
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS mandatory (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                name TEXT NOT NULL,
                amount REAL NOT NULL
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS expenses (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                spent_date TEXT NOT NULL,
                amount REAL NOT NULL,
                category TEXT NOT NULL,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS savings_topups (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                topup_date TEXT NOT NULL,
                amount REAL NOT NULL,
                comment TEXT NOT NULL,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP
            )
        """)
        # снимки дневного лимита: прошедшие дни не «пересчитываются» после смены настроек
        conn.execute("""
            CREATE TABLE IF NOT EXISTS day_limits (
                user_id INTEGER NOT NULL,
                day TEXT NOT NULL,
                daily_limit REAL NOT NULL,
                PRIMARY KEY (user_id, day)
            )
        """)
        # связь «снятие из копилки» <-> трата, чтобы удалять их вместе
        cols = [r["name"] for r in conn.execute("PRAGMA table_info(savings_topups)")]
        if "expense_id" not in cols:
            conn.execute("ALTER TABLE savings_topups ADD COLUMN expense_id INTEGER")
        # миграция: снятия из копилки больше не дублируются «тратой дня» (она ломала лимит и банк сдачи)
        conn.execute(
            "DELETE FROM expenses WHERE id IN "
            "(SELECT expense_id FROM savings_topups WHERE expense_id IS NOT NULL)"
        )
        conn.execute("UPDATE savings_topups SET expense_id = NULL WHERE expense_id IS NOT NULL")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_exp_user_date ON expenses(user_id, spent_date)")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_top_user_date ON savings_topups(user_id, topup_date)")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_mand_user ON mandatory(user_id)")


def get_setting(user_id: int, key: str, default=None):
    with db() as conn:
        row = conn.execute(
            "SELECT value FROM settings WHERE user_id = ? AND key = ?",
            (user_id, key),
        ).fetchone()
        return row["value"] if row else default


LIMIT_KEYS = {"daily_income", "income_mode", "save", "custom_days_in_month"}


def set_setting(user_id: int, key: str, value):
    if key in LIMIT_KEYS:
        freeze_past_limits(user_id)  # прошлые дни остаются при старом лимите
    with db() as conn:
        conn.execute(
            "INSERT INTO settings (user_id, key, value) VALUES (?, ?, ?) "
            "ON CONFLICT(user_id, key) DO UPDATE SET value = excluded.value",
            (user_id, key, str(value)),
        )


# --- Вспомогательные функции ---
def today_date() -> datetime.date:
    return datetime.datetime.now(TZ).date()


def fmt_money(val: float) -> str:
    return f"{val:,.0f}".replace(",", " ") + " ₽"


def fmt_num(val: float) -> str:
    return f"{int(round(val)):,}".replace(",", " ")


def esc(s) -> str:
    """Экранирование пользовательского текста для parse_mode=Markdown."""
    return escape_markdown(str(s), version=1)


def cap(s: str) -> str:
    s = s.strip()
    return s[:1].upper() + s[1:]


_CURRENCY_TAIL_RE = re.compile(r"(?:₽|руб\.?|р\.?|rub)$", re.IGNORECASE)
_AMOUNT_RE = re.compile(r"\d+(?:\.\d+)?")
_THOUSANDS_COMMA_RE = re.compile(r"\d{1,3}(?:,\d{3})+")


def parse_amount(s: str, allow_zero: bool = False):
    """Число из строки или None: «1500», «1 500», «1,000», «300₽», «1500р», «10,5».
    Отсекает inf/nan, экспоненту, отрицательные и слишком большие."""
    if not isinstance(s, str):
        return None
    t = s.replace("\u00a0", "").replace(" ", "").replace("₽", "")
    t = _CURRENCY_TAIL_RE.sub("", t)
    if _THOUSANDS_COMMA_RE.fullmatch(t):
        t = t.replace(",", "")
    else:
        t = t.replace(",", ".")
    if not _AMOUNT_RE.fullmatch(t):
        return None
    v = float(t)
    if not math.isfinite(v) or v > MAX_AMOUNT:
        return None
    if v == 0 and not allow_zero:
        return None
    return v


def merge_thousands(tokens):
    """«1 500 такси» -> «1500 такси»: склеивает группы цифр, разделённые пробелом."""
    out = []
    for tok in tokens:
        if (
            out
            and re.fullmatch(r"\d{3}(?:[.,]\d{1,2})?", tok)
            and re.fullmatch(r"\d{1,3}(?:\d{3})*", out[-1])
        ):
            out[-1] += tok
        else:
            out.append(tok)
    return out


def parse_time(text: str):
    m = re.match(r"^(\d{1,2}):(\d{2})$", text.strip())
    if not m:
        return None
    h, mn = int(m.group(1)), int(m.group(2))
    if not (0 <= h <= 23 and 0 <= mn <= 59):
        return None
    return h, mn


def save_reminder_time(user_id: int, h: int, mn: int):
    set_setting(user_id, "reminder_time", f"{h:02d}:{mn:02d}")
    now = datetime.datetime.now(TZ)
    # если время на сегодня уже прошло, отчёт придёт только завтра
    if now.hour * 60 + now.minute >= h * 60 + mn:
        set_setting(user_id, "last_report", today_date().isoformat())


def ensure_start_day(user_id: int):
    if not get_setting(user_id, "start_date"):
        set_setting(user_id, "start_date", today_date().isoformat())


def get_period_start(user_id: int, ref_date: datetime.date) -> datetime.date:
    """С какого дня считаем месяц: начало месяца или день старта, если он позже."""
    ensure_start_day(user_id)
    try:
        start_d = datetime.date.fromisoformat(get_setting(user_id, "start_date"))
    except (ValueError, TypeError):
        start_d = ref_date
    return max(start_d, ref_date.replace(day=1))


# --- Бюджет ---
def get_budget_info(user_id: int, ref_date: datetime.date = None):
    if ref_date is None:
        ref_date = today_date()

    real_days = calendar.monthrange(ref_date.year, ref_date.month)[1]
    custom_days = get_setting(user_id, "custom_days_in_month")
    days_in_month = int(custom_days) if custom_days and str(custom_days).isdigit() else real_days

    income_mode = get_setting(user_id, "income_mode", "daily")
    raw_income = float(get_setting(user_id, "daily_income", 1000) or 1000)

    if income_mode == "monthly":
        income = raw_income
        display_income_text = f"{fmt_money(raw_income)} 1 раз в месяц"
    elif income_mode == "twice_monthly":
        income = raw_income * 2
        display_income_text = f"{fmt_money(raw_income)} 2 раза в месяц"
    elif income_mode == "weekly":
        income = raw_income * (days_in_month / 7)
        display_income_text = f"{fmt_money(raw_income)} раз в неделю"
    else:
        income_mode = "daily"
        income = raw_income * days_in_month
        display_income_text = f"по {fmt_money(raw_income)} в день"

    save_advance = float(get_setting(user_id, "save", 0) or 0)

    month_prefix = f"{ref_date.year:04d}-{ref_date.month:02d}-"
    with db() as conn:
        mand_rows = conn.execute(
            "SELECT id, name, amount FROM mandatory WHERE user_id = ? ORDER BY id",
            (user_id,),
        ).fetchall()
        topup_row = conn.execute(
            "SELECT SUM(amount) AS total FROM savings_topups "
            "WHERE user_id = ? AND topup_date LIKE ? AND amount > 0",
            (user_id, month_prefix + "%"),
        ).fetchone()
        all_topups_row = conn.execute(
            "SELECT SUM(amount) AS total FROM savings_topups WHERE user_id = ? AND topup_date <= ?",
            (user_id, ref_date.isoformat()),
        ).fetchone()

    topups_total = topup_row["total"] if topup_row and topup_row["total"] else 0.0
    all_topups = all_topups_row["total"] if all_topups_row and all_topups_row["total"] else 0.0
    # накопленная копилка: плановые отчисления за каждый месяц с начала учёта + все пополнения/снятия
    try:
        start_d = datetime.date.fromisoformat(get_setting(user_id, "start_date") or "")
    except ValueError:
        start_d = ref_date
    n_months = max(0, (ref_date.year - start_d.year) * 12 + ref_date.month - start_d.month + 1)
    piggy_total = save_advance * n_months + all_topups
    total_save_effective = save_advance + topups_total

    mandatory_total = sum(r["amount"] for r in mand_rows)
    available = income - mandatory_total - total_save_effective
    limit_base = available / days_in_month if days_in_month else 0.0

    with db() as conn:
        snap_rows = conn.execute(
            "SELECT day, daily_limit FROM day_limits WHERE user_id = ? AND day LIKE ?",
            (user_id, month_prefix + "%"),
        ).fetchall()
    limit_snaps = {int(r["day"][-2:]): r["daily_limit"] for r in snap_rows}
    first_counted_day = 1
    try:
        sd = datetime.date.fromisoformat(get_setting(user_id, "start_date") or "")
        if (sd.year, sd.month) == (ref_date.year, ref_date.month):
            first_counted_day = sd.day
    except ValueError:
        pass
    series = _month_limit_series(limit_base, days_in_month, limit_snaps, first_counted_day, ref_date.day)
    daily_limit = series.get(ref_date.day, max(0.0, limit_base))

    return {
        "raw_income": raw_income,
        "income_mode": income_mode,
        "display_income_text": display_income_text,
        "income": income,
        "save": save_advance,
        "topups_total": topups_total,
        "total_save_effective": total_save_effective,
        "piggy_total": piggy_total,
        "mandatory_list": mand_rows,
        "mandatory_total": mandatory_total,
        "days_in_month": days_in_month,
        "real_days": real_days,
        "is_custom_days": bool(custom_days),
        "available": available,
        "daily_limit": max(0.0, daily_limit),
        "limit_base": limit_base,
        "limit_snaps": limit_snaps,
        "first_counted_day": first_counted_day,
    }


def _month_limit_series(base: float, n: int, snaps: dict, first_day: int, last_day: int) -> dict:
    """Лимит по каждому дню месяца. Дни со снимком держат старый лимит, остальные
    делят оставшийся бюджет поровну: так смена настроек посреди месяца не меняет прошлое,
    а сумма лимитов за месяц по-прежнему равна «доход − обязательные − копилка»."""
    out, delta = {}, 0.0
    for dn in range(first_day, last_day + 1):
        if dn in snaps:
            lim = snaps[dn]
        else:
            lim = max(0.0, base - delta / max(1, n - (dn - 1)))
        out[dn] = lim
        delta += lim - base
    return out


def limits_between(user_id: int, start: datetime.date, end: datetime.date) -> dict:
    """{дата: лимит} для диапазона дат (дни до начала учёта в результат не попадают)."""
    out = {}
    cur = start.replace(day=1)
    while cur <= end:
        last_dn = calendar.monthrange(cur.year, cur.month)[1]
        info = get_budget_info(user_id, cur)
        series = _month_limit_series(
            info["limit_base"], info["days_in_month"], info["limit_snaps"],
            info["first_counted_day"], last_dn,
        )
        for dn, lim in series.items():
            out[datetime.date(cur.year, cur.month, dn)] = lim
        cur = (cur.replace(day=28) + datetime.timedelta(days=4)).replace(day=1)
    return {d: v for d, v in out.items() if start <= d <= end}


def freeze_past_limits(user_id: int):
    """Фиксирует лимит всех прошедших дней. Вызывается ПЕРЕД любым изменением,
    влияющим на лимит (доход, копилка, обязательные расходы, число дней)."""
    raw = get_setting(user_id, "start_date")
    if not raw or get_setting(user_id, "daily_income") is None:
        return  # бюджет ещё не настроен — фиксировать нечего (иначе заморозится лимит по умолчанию)
    try:
        start = datetime.date.fromisoformat(raw)
    except ValueError:
        return
    yesterday = today_date() - datetime.timedelta(days=1)
    if start > yesterday:
        return
    lims = limits_between(user_id, start, yesterday)
    with db() as conn:
        conn.executemany(
            "INSERT OR IGNORE INTO day_limits (user_id, day, daily_limit) VALUES (?, ?, ?)",
            [(user_id, d.isoformat(), v) for d, v in lims.items()],
        )


def month_limits(user_id: int, today: datetime.date, start: datetime.date) -> dict:
    """{номер дня: лимит} от начала учёта до конца текущего месяца."""
    real_days = calendar.monthrange(today.year, today.month)[1]
    lims = limits_between(user_id, start, datetime.date(today.year, today.month, real_days))
    return {d.day: v for d, v in lims.items()}


def upsert_mandatory(user_id: int, name: str, amount: float) -> bool:
    """Добавляет обязательный расход; расход с тем же названием обновляет. True, если обновлён."""
    freeze_past_limits(user_id)
    with db() as conn:
        rows = conn.execute("SELECT id, name FROM mandatory WHERE user_id = ?", (user_id,)).fetchall()
        for r in rows:
            if r["name"].strip().lower() == name.strip().lower():
                conn.execute("UPDATE mandatory SET amount = ? WHERE id = ?", (amount, r["id"]))
                return True
        conn.execute(
            "INSERT INTO mandatory (user_id, name, amount) VALUES (?, ?, ?)",
            (user_id, name, amount),
        )
    return False


def calc_change_pool(user_id: int, through_date: datetime.date) -> float:
    """Банк сдачи: лимит за все дни периода по through_date включительно минус траты."""
    period_start = get_period_start(user_id, through_date)
    if period_start > through_date:
        return 0.0
    total_limit = sum(limits_between(user_id, period_start, through_date).values())
    with db() as conn:
        row = conn.execute(
            "SELECT SUM(amount) AS total FROM expenses "
            "WHERE user_id = ? AND spent_date >= ? AND spent_date <= ?",
            (user_id, period_start.isoformat(), through_date.isoformat()),
        ).fetchone()
    spent = row["total"] or 0.0
    return total_limit - spent


def delete_expense(user_id: int, exp_id: int):
    """Удаляет трату вместе со связанным снятием из копилки. Возвращает дату траты или None."""
    with db() as conn:
        row = conn.execute(
            "SELECT spent_date FROM expenses WHERE id = ? AND user_id = ?",
            (exp_id, user_id),
        ).fetchone()
        if not row:
            return None
        conn.execute(
            "DELETE FROM savings_topups WHERE expense_id = ? AND user_id = ?",
            (exp_id, user_id),
        )
        conn.execute("DELETE FROM expenses WHERE id = ? AND user_id = ?", (exp_id, user_id))
        return row["spent_date"]


# --- Экспорт в Excel ---
def update_excel_file(user_id: int, target_month: datetime.date = None) -> io.BytesIO:
    if target_month is None:
        target_month = today_date()

    first_day = datetime.date(target_month.year, target_month.month, 1)
    info = get_budget_info(user_id, target_month)
    real_days = info["real_days"]

    if os.path.exists(TEMPLATE_PATH):
        wb = openpyxl.load_workbook(TEMPLATE_PATH)
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Бюджет"

    ws["B4"] = datetime.datetime(first_day.year, first_day.month, 1)
    ws["B5"] = info["income"]
    ws["B20"] = info["total_save_effective"]
    if info["is_custom_days"]:
        ws["B23"] = info["days_in_month"]  # иначе Excel делит на реальные дни и расходится с ботом

    for row_idx in range(8, 18):
        ws[f"A{row_idx}"] = None
        ws[f"B{row_idx}"] = None

    mand = list(info["mandatory_list"])
    if len(mand) > 10:
        shown = mand[:9]
        rest = mand[9:]
        for idx, item in enumerate(shown):
            ws[f"A{8 + idx}"] = item["name"]
            ws[f"B{8 + idx}"] = item["amount"]
        ws["A17"] = f"Остальные обязательные ({len(rest)} шт.)"
        ws["B17"] = sum(r["amount"] for r in rest)
    else:
        for idx, item in enumerate(mand):
            ws[f"A{8 + idx}"] = item["name"]
            ws[f"B{8 + idx}"] = item["amount"]

    for row_idx in range(6, 37):
        ws[f"G{row_idx}"] = None
        ws[f"H{row_idx}"] = None

    month_prefix = f"{first_day.year:04d}-{first_day.month:02d}-"
    with db() as conn:
        rows = conn.execute(
            "SELECT spent_date, amount, category FROM expenses "
            "WHERE user_id = ? AND spent_date LIKE ? ORDER BY spent_date, id",
            (user_id, month_prefix + "%"),
        ).fetchall()

    by_day = {}
    for r in rows:
        d_num = int(r["spent_date"].split("-")[2])
        by_day.setdefault(d_num, []).append((r["amount"], r["category"]))

    lims = limits_between(user_id, first_day, datetime.date(first_day.year, first_day.month, real_days))
    t_today = today_date()
    try:
        start_d = datetime.date.fromisoformat(get_setting(user_id, "start_date", t_today.isoformat()))
    except ValueError:
        start_d = t_today

    for day_num in range(1, real_days + 1):
        cur_date = datetime.date(first_day.year, first_day.month, day_num)
        row_idx = 5 + day_num
        if day_num <= info["days_in_month"]:
            ws[f"F{row_idx}"] = round(lims.get(cur_date, info["daily_limit"]), 2)
        if day_num in by_day:
            items = by_day[day_num]
            ws[f"G{row_idx}"] = sum(x[0] for x in items)
            ws[f"H{row_idx}"] = " | ".join(
                f"{cat}: {int(amt) if float(amt).is_integer() else amt}₽" for amt, cat in items
            )
        elif start_d <= cur_date < t_today:
            ws[f"G{row_idx}"] = 0
            ws[f"H{row_idx}"] = "Без трат"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# --- Клавиатуры ---
MAIN_KB = ReplyKeyboardMarkup(
    [
        ["📊 Сегодня", "📅 Месяц", "🏦 В копилку"],
        ["💸 Из копилки", "📈 График", "📥 Скачать Excel"],
        ["⚙️ Настройки", "↩️ Отменить последнюю"],
    ],
    resize_keyboard=True,
)
MENU_TEXTS = {b for row in MAIN_KB.keyboard for b in (x.text if hasattr(x, "text") else x for x in row)}


def mode_markup(prefix: str):
    return InlineKeyboardMarkup(
        [[InlineKeyboardButton(label, callback_data=prefix + key)] for key, label in INCOME_MODES.items()]
    )


def chart_kb():
    return InlineKeyboardMarkup([
        [
            InlineKeyboardButton("📊 Лимит и банк", callback_data="chart_overview"),
            InlineKeyboardButton("🗓 Календарь", callback_data="chart_calendar"),
        ],
        [
            InlineKeyboardButton("🔮 Прогноз", callback_data="chart_forecast"),
            InlineKeyboardButton("🍩 Структура", callback_data="chart_structure"),
        ],
        [InlineKeyboardButton("💧 Итог месяца", callback_data="chart_waterfall")],
    ])


def settings_inline_kb(user_id: int):
    info = get_budget_info(user_id)
    reminder_time = get_setting(user_id, "reminder_time", "21:00")
    days_label = (
        f"{info['days_in_month']} дн. (авто)"
        if not info["is_custom_days"]
        else f"{info['days_in_month']} дн. (ручное)"
    )
    chart_on = get_setting(user_id, "report_chart", "1") == "1"
    mode_label = INCOME_MODE_SHORT.get(info["income_mode"], info["income_mode"])
    return InlineKeyboardMarkup([
        [InlineKeyboardButton(f"🏦 Доход: {fmt_money(info['raw_income'])} ({mode_label})", callback_data="set_daily_income")],
        [InlineKeyboardButton("🔄 Изменить график получения дохода", callback_data="set_income_mode")],
        [InlineKeyboardButton(f"📅 Дней в месяце: {days_label}", callback_data="set_days_count")],
        [InlineKeyboardButton(f"⏰ Время отчёта, {TZ_LABEL} (сейчас {reminder_time})", callback_data="set_reminder_time")],
        [InlineKeyboardButton(f"📈 График в вечернем отчёте: {'вкл' if chart_on else 'выкл'}", callback_data="toggle_report_chart")],
        [InlineKeyboardButton("💰 Плановая копилка", callback_data="set_save")],
        [InlineKeyboardButton("🏦 История копилки", callback_data="piggy_history")],
        [
            InlineKeyboardButton("➕ Добавить обяз. расход", callback_data="add_mand"),
            InlineKeyboardButton("🗑 Удалить обяз. расход", callback_data="del_mand_menu"),
        ],
        [InlineKeyboardButton("🔄 Запустить мастера настройки", callback_data="restart_wizard")],
        [InlineKeyboardButton("🧹 Сбросить текущий ввод", callback_data="restart_bot_action")],
        [InlineKeyboardButton("⚠️ Полный сброс всех данных", callback_data="full_reset_action")],
    ])


def build_settings_text(user_id: int) -> str:
    ensure_start_day(user_id)
    info = get_budget_info(user_id)
    mand_lines = [f"  • {esc(r['name'])}: {fmt_money(r['amount'])}" for r in info["mandatory_list"]]
    mand_block = "\n".join(mand_lines) if mand_lines else "  (пока не добавлены)"

    topup_str = f" (+{fmt_money(info['topups_total'])} доп.)" if info["topups_total"] > 0 else ""
    reminder_time = get_setting(user_id, "reminder_time", "21:00")
    days_mode = "авто по календарю" if not info["is_custom_days"] else "ручной выбор"

    text = (
        "⚙️ *Настройки бюджета (🏦 Умная модель)*\n\n"
        f"🏦 *Доход:* {fmt_money(info['raw_income'])} ({info['display_income_text']})\n"
        f"📅 *Дней в месяце:* {info['days_in_month']} ({days_mode})\n"
        f"⏰ *Вечерний авто-отчёт:* в {reminder_time} ({TZ_LABEL})\n"
        f"🏠 *Обязательные расходы ({fmt_money(info['mandatory_total'])}):*\n{mand_block}\n\n"
        f"💰 *Плановая копилка:* {fmt_money(info['save'])}{topup_str}\n"
        f"🏦 *Копилка (всего накоплено):* {fmt_money(info['piggy_total'])}\n"
        f"🎯 *Чистый лимит на день:* *{fmt_money(info['daily_limit'])}*\n"
    )
    if info["available"] < 0:
        text += (
            f"\n⚠️ Обязательные расходы и копилка больше дохода на "
            f"*{fmt_money(-info['available'])}*, поэтому лимит на день равен 0.\n"
        )
    return text + "\n👇 Нажмите нужную кнопку для изменения:"


def build_day_report(user_id: int, target_date: datetime.date):
    info = get_budget_info(user_id, target_date)
    with db() as conn:
        rows = conn.execute(
            "SELECT id, amount, category FROM expenses WHERE user_id = ? AND spent_date = ? ORDER BY id",
            (user_id, target_date.isoformat()),
        ).fetchall()

    spent = sum(r["amount"] for r in rows)
    rem = info["daily_limit"] - spent
    change_pool = calc_change_pool(user_id, target_date)

    date_str = target_date.strftime("%d.%m.%Y")
    title = (
        f"📊 *Траты за сегодня ({date_str})*"
        if target_date == today_date()
        else f"📊 *Траты за {date_str}*"
    )

    if rows:
        exp_text = "\n".join(
            f"{i + 1}. *{fmt_money(r['amount'])}* — {esc(r['category'])}" for i, r in enumerate(rows)
        )
    else:
        exp_text = "_Пока нет записанных трат_"

    if rem < 0:
        status_line = f"⚠️ *ЛИМИТ ПРЕВЫШЕН на {fmt_money(-rem)}!*"
        if change_pool >= 0:
            advice = "💡 Перерасход покрыт накопленной сдачей."
        else:
            advice = "💡 Банк сдачи ушёл в минус. Попробуйте завтра потратить меньше."
    else:
        status_line = f"✅ *Остаток сдачи за день:* +{fmt_money(rem)}"
        advice = f"💡 Сдача (+{fmt_money(rem)}) добавлена в ваш банк сдачи."

    base_daily_income = info["income"] / info["days_in_month"] if info["days_in_month"] else 0.0

    text = (
        f"{title}\n\n"
        f"{exp_text}\n\n"
        f"🏦 *Доход за день (расчётный):* +{fmt_money(base_daily_income)}\n"
        f"💸 *Потрачено за день:* {fmt_money(spent)} из {fmt_money(info['daily_limit'])}\n"
        f"{status_line}\n\n"
        f"{advice}\n\n"
        f"🪙 *Банк сдачи:* {fmt_money(change_pool)}\n"
        f"💰 *Копилка (всего):* {fmt_money(info['piggy_total'])}"
    )
    if info["available"] < 0:
        text += (
            f"\n\n⚠️ Обязательные расходы и копилка больше дохода на "
            f"{fmt_money(-info['available'])}. Лимит на день: 0 ₽."
        )

    buttons = []
    for r in rows:
        btn_title = f"❌ {r['category']} ({int(r['amount'])}₽)"
        if len(btn_title) > 35:
            btn_title = btn_title[:32] + "..."
        buttons.append(InlineKeyboardButton(btn_title, callback_data=f"delexp_{r['id']}"))

    kb_rows = [buttons[i:i + 2] for i in range(0, len(buttons), 2)]
    markup = InlineKeyboardMarkup(kb_rows) if kb_rows else None
    return text, markup


# --- Разбор сообщений ---
DATE_WORDS = {"сегодня": 0, "вчера": -1, "позавчера": -2}
DATE_FULL_RE = re.compile(r"^(\d{1,2})\.(\d{1,2})\.(\d{2,4})$")
DATE_SHORT_RE = re.compile(r"^(\d{1,2})\.(\d{1,2})$")


def _make_date(d: int, m: int, y: int):
    if y < 100:
        y += 2000
    try:
        return datetime.date(y, m, d)
    except ValueError:
        return None


def parse_expense_message(text: str):
    """«450 продукты», «такси 300 вчера», «450 кофе 28.10». Возвращает (сумма, категория, дата)."""
    tokens = merge_thousands(text.strip().split())
    if not tokens:
        return None

    base = today_date()
    target = base
    rest = []

    for tok in tokens:
        low = tok.lower()
        if low in DATE_WORDS:
            target = base + datetime.timedelta(days=DATE_WORDS[low])
            continue
        m = DATE_FULL_RE.match(tok)
        if m:
            dt = _make_date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
            if dt:
                target = dt
                continue
        rest.append(tok)

    def short_date(tok):
        m = DATE_SHORT_RE.match(tok)
        if not m:
            return None
        dt = _make_date(int(m.group(1)), int(m.group(2)), base.year)
        if dt and dt > base:
            dt = _make_date(int(m.group(1)), int(m.group(2)), base.year - 1)
        # «1.5» в «2 молоко 1.5» — скорее число, чем дата многомесячной давности
        if dt and 0 <= (base - dt).days <= 45:
            return dt
        return None

    # «10.5» это дата, только если в сообщении есть ещё и обычная сумма; иначе это 10,5 ₽
    if any(parse_amount(t) is not None and short_date(t) is None for t in rest):
        remaining = []
        for tok in rest:
            sd = short_date(tok)
            if sd:
                target = sd
            else:
                remaining.append(tok)
        rest = remaining

    if not rest:
        return None

    first = parse_amount(rest[0])
    last = parse_amount(rest[-1]) if len(rest) > 1 else None

    if first is not None:
        amount, cat_tokens = first, rest[1:]
    elif last is not None:
        amount, cat_tokens = last, rest[:-1]
    else:
        return None

    category = cap(" ".join(cat_tokens)) or "Разное"
    return amount, category[:60], target


def parse_label_amount(text: str):
    """«Аренда 15000». Возвращает (сумма, название) без разбора дат."""
    tokens = merge_thousands(text.strip().split())
    if not tokens:
        return None
    first = parse_amount(tokens[0])
    last = parse_amount(tokens[-1]) if len(tokens) > 1 else None
    if first is not None:
        amount, name_tokens = first, tokens[1:]
    elif last is not None:
        amount, name_tokens = last, tokens[:-1]
    else:
        return None
    return amount, (cap(" ".join(name_tokens)) or "Разное")[:60]


def parse_amount_comment(text: str, default_comment: str):
    """«500 Подработка». Возвращает (сумма, комментарий) или None."""
    tokens = merge_thousands(text.strip().split())
    if not tokens:
        return None
    amount = parse_amount(tokens[0])
    if amount is None:
        return None
    comment = cap(" ".join(tokens[1:])) if len(tokens) > 1 else default_comment
    return amount, comment[:60]


# --- Графики ---
def collect_month(user_id: int):
    today = today_date()
    info = get_budget_info(user_id, today)
    start = get_period_start(user_id, today)
    prefix = f"{today.year:04d}-{today.month:02d}-"
    with db() as conn:
        day_rows = conn.execute(
            "SELECT spent_date, SUM(amount) AS total FROM expenses "
            "WHERE user_id = ? AND spent_date LIKE ? GROUP BY spent_date",
            (user_id, prefix + "%"),
        ).fetchall()
        cat_rows = conn.execute(
            "SELECT category, SUM(amount) AS total FROM expenses "
            "WHERE user_id = ? AND spent_date LIKE ? GROUP BY category ORDER BY total DESC",
            (user_id, prefix + "%"),
        ).fetchall()
    by_day = {int(r["spent_date"].split("-")[2]): r["total"] for r in day_rows}
    return today, info, start, by_day, cat_rows


def _to_buf(fig) -> io.BytesIO:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=130, bbox_inches="tight", facecolor="white")
    buf.seek(0)
    return buf


def _style(ax):
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    ax.grid(axis="y", linestyle=":", alpha=0.5)
    ax.set_axisbelow(True)
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: fmt_num(v)))


def _xticks(ax, days):
    step = max(1, len(days) // 12)
    ax.set_xticks(days[::step])


def chart_overview(user_id: int) -> io.BytesIO:
    today, info, start, by_day, _ = collect_month(user_id)
    limit = info["daily_limit"]
    days = list(range(start.day, today.day + 1)) or [today.day]
    spent = [by_day.get(d, 0.0) for d in days]
    lm = month_limits(user_id, today, start)
    lim_d = [lm.get(d, limit) for d in days]

    fig = Figure(figsize=(8, 6.4))
    ax1, ax2 = fig.subplots(2, 1)

    ax1.bar(days, spent, color=[RED if s > l else BLUE for s, l in zip(spent, lim_d)], width=0.7)
    ax1.step(days, lim_d, where="mid", color=GREEN, linestyle="--", linewidth=2)
    ax1.set_title("Траты по дням и дневной лимит, ₽", fontsize=11, fontweight="bold", loc="left")
    ax1.legend(
        handles=[
            Patch(facecolor=BLUE, label="В лимите"),
            Patch(facecolor=RED, label="Перерасход"),
            Line2D([0], [0], color=GREEN, linestyle="--", linewidth=2, label=f"Лимит {fmt_num(limit)} ₽"),
        ],
        frameon=False,
        fontsize=9,
        ncol=3,
        loc="upper left",
        bbox_to_anchor=(0, 1.0),
    )
    _style(ax1)
    _xticks(ax1, days)

    cum, run = [], 0.0
    for s, l in zip(spent, lim_d):
        run += l - s
        cum.append(run)
    ax2.axhline(0, color=GRAY, linewidth=1)
    ax2.fill_between(days, cum, 0, where=[c >= 0 for c in cum], color=GREEN, alpha=0.18, interpolate=True)
    ax2.fill_between(days, cum, 0, where=[c < 0 for c in cum], color=RED, alpha=0.18, interpolate=True)
    ax2.plot(days, cum, color=BLUE, linewidth=2, marker="o", markersize=3)
    sign = "+" if cum[-1] >= 0 else "−"
    ax2.set_title(
        f"Банк сдачи нарастающим итогом: {sign}{fmt_num(abs(cum[-1]))} ₽",
        fontsize=11, fontweight="bold", loc="left",
    )
    _style(ax2)
    _xticks(ax2, days)
    ax2.set_xlabel("День месяца", color=GRAY)

    fig.tight_layout()
    return _to_buf(fig)


def chart_calendar(user_id: int) -> io.BytesIO:
    today, info, start, by_day, _ = collect_month(user_id)
    limit = info["daily_limit"]
    weeks = calendar.monthcalendar(today.year, today.month)
    n = len(weeks)
    lm = month_limits(user_id, today, start)

    fig = Figure(figsize=(7.6, 2.0 + 0.95 * n))
    ax = fig.add_subplot(111)
    ax.set_xlim(0, 7)
    ax.set_ylim(n + 0.8, -0.8)
    ax.axis("off")

    for c, name in enumerate(["Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Вс"]):
        ax.text(c + 0.5, -0.35, name, ha="center", va="center", fontsize=10, color=GRAY)

    ok_days = active_days = 0
    saved = 0.0
    for r, week in enumerate(weeks):
        for c, d in enumerate(week):
            if d == 0:
                continue
            dt = datetime.date(today.year, today.month, d)
            label = ""
            if start <= dt <= today:
                lim_d = lm.get(d, limit)
                diff = lim_d - by_day.get(d, 0.0)
                active_days += 1
                saved += diff
                if diff >= 0:
                    ok_days += 1
                intensity = min(abs(diff) / lim_d, 1.0) if lim_d > 0 else 1.0
                face = to_rgba(GREEN_D if diff >= 0 else RED, 0.22 + 0.55 * intensity)
                label = f"{'+' if diff >= 0 else '−'}{fmt_num(abs(diff))}"
            else:
                face = to_rgba("#9ca3af", 0.12)
            is_today = dt == today
            ax.add_patch(Rectangle(
                (c + 0.04, r + 0.04), 0.92, 0.92,
                facecolor=face,
                edgecolor=INK if is_today else "none",
                linewidth=1.8 if is_today else 0,
            ))
            ax.text(c + 0.12, r + 0.22, str(d), fontsize=9, color="#374151", va="center")
            if label:
                ax.text(c + 0.5, r + 0.63, label, ha="center", va="center",
                        fontsize=10, fontweight="bold", color=INK)

    streak, cur = 0, today
    while cur >= start and by_day.get(cur.day, 0.0) <= lm.get(cur.day, limit):
        streak += 1
        cur -= datetime.timedelta(days=1)

    sign = "+" if saved >= 0 else "−"
    ax.text(
        3.5, n + 0.4,
        f"Серия в лимите: {streak} дн.    В лимите: {ok_days} из {active_days} дн.    "
        f"Итог: {sign}{fmt_num(abs(saved))} ₽",
        ha="center", va="center", fontsize=10, color=INK,
    )
    fig.suptitle(
        f"{MONTHS_RU[today.month - 1]} {today.year}: отклонение от лимита {fmt_num(limit)} ₽/день",
        fontsize=12, fontweight="bold",
    )
    return _to_buf(fig)


def chart_forecast(user_id: int) -> io.BytesIO:
    today, info, start, by_day, _ = collect_month(user_id)
    limit = info["daily_limit"]
    horizon = max(info["days_in_month"], today.day)
    days_all = list(range(start.day, horizon + 1)) or [today.day]
    lm = month_limits(user_id, today, start)
    limit_line, acc = [], 0.0
    for d in days_all:
        acc += lm.get(d, limit)
        limit_line.append(acc)

    fact_days = list(range(start.day, today.day + 1)) or [today.day]
    cum, run = [], 0.0
    for d in fact_days:
        run += by_day.get(d, 0.0)
        cum.append(run)
    spent_total = cum[-1]
    n_elapsed = max(1, today.day - start.day + 1)
    avg = spent_total / n_elapsed

    fc_days = list(range(today.day, horizon + 1))
    fc = [spent_total + avg * (d - today.day) for d in fc_days]
    rest = limit_line[-1] - fc[-1]
    fc_color = BLUE if rest >= 0 else RED

    fig = Figure(figsize=(8, 4.8))
    ax = fig.add_subplot(111)
    ax.plot(days_all, limit_line, color=GRAY, linewidth=2)
    ax.fill_between(fact_days, cum, color=BLUE, alpha=0.10)
    ax.plot(fact_days, cum, color=BLUE, linewidth=2)
    ax.plot(fc_days, fc, color=fc_color, linewidth=2, linestyle="--")
    ax.scatter([today.day], [spent_total], s=50, color=BLUE, zorder=5, edgecolor="white", linewidth=1.5)
    if len(fc_days) > 1:
        ax.scatter([fc_days[-1]], [fc[-1]], s=50, color=fc_color, zorder=5, edgecolor="white", linewidth=1.5)
    ax.set_ylim(bottom=0)
    ax.set_xlim(days_all[0] - 0.5, days_all[-1] + 0.5)
    _style(ax)
    ax.set_xlabel("День месяца", color=GRAY)
    ax.legend(
        handles=[
            Line2D([0], [0], color=BLUE, linewidth=2, label="Факт"),
            Line2D([0], [0], color=fc_color, linewidth=2, linestyle="--", label="Прогноз"),
            Line2D([0], [0], color=GRAY, linewidth=2, label="Допустимо"),
        ],
        frameon=False, fontsize=9, ncol=3, loc="upper left",
    )
    verdict = "останется ≈" if rest >= 0 else "перерасход ≈"
    ax.set_title(
        f"Прогноз: при темпе {fmt_num(avg)} ₽/день к {horizon}-му {verdict} {fmt_num(abs(rest))} ₽",
        fontsize=11, fontweight="bold", loc="left",
    )
    fig.tight_layout()
    return _to_buf(fig)


def chart_structure(user_id: int) -> io.BytesIO:
    today, info, start, by_day, cats = collect_month(user_id)
    fig = Figure(figsize=(9.5, 4.6))
    ax1, ax2 = fig.subplots(1, 2, gridspec_kw={"width_ratios": [1, 1.25]})

    parts = [
        ("Обязательные", info["mandatory_total"], BLUE),
        ("Копилка", max(0.0, info["total_save_effective"]), GREEN),
        ("Свободные деньги", max(0.0, info["available"]), ORANGE),
    ]
    parts = [p for p in parts if p[1] > 0]
    if parts:
        ax1.pie(
            [p[1] for p in parts],
            colors=[p[2] for p in parts],
            startangle=90,
            counterclock=False,
            wedgeprops={"width": 0.34, "edgecolor": "white", "linewidth": 2},
        )
        ax1.text(0, 0, f"{fmt_num(info['daily_limit'])} ₽\nв день",
                 ha="center", va="center", fontsize=13, fontweight="bold")
        ax1.legend(
            handles=[Patch(facecolor=c, label=f"{n}: {fmt_num(v)} ₽") for n, v, c in parts],
            loc="upper center", bbox_to_anchor=(0.5, 0.02), frameon=False, fontsize=9,
        )
    else:
        ax1.axis("off")
        ax1.text(0.5, 0.5, "Нет данных", ha="center", va="center", transform=ax1.transAxes)
    title1 = "Структура дохода"
    if info["available"] < 0:
        title1 += f" (дефицит {fmt_num(-info['available'])} ₽)"
    ax1.set_title(title1, fontsize=11, fontweight="bold")

    top = [(r["category"], r["total"]) for r in cats[:7]]
    other = sum(r["total"] for r in cats[7:])
    if other > 0:
        top.append(("Прочее", other))
    if top:
        names = [n if len(n) <= 22 else n[:21] + "…" for n, _ in top]
        vals = [v for _, v in top]
        ax2.barh(range(len(vals)), vals, color=BLUE, height=0.6)
        ax2.set_yticks(range(len(vals)))
        ax2.set_yticklabels(names, fontsize=10)
        ax2.invert_yaxis()
        ax2.set_xlim(0, max(vals) * 1.2)
        for i, v in enumerate(vals):
            ax2.text(v + max(vals) * 0.015, i, fmt_num(v), va="center", fontsize=9, color=INK)
        for side in ("top", "right"):
            ax2.spines[side].set_visible(False)
        ax2.xaxis.set_major_formatter(FuncFormatter(lambda v, _: fmt_num(v)))
        ax2.grid(axis="x", linestyle=":", alpha=0.5)
        ax2.set_axisbelow(True)
    else:
        ax2.axis("off")
        ax2.text(0.5, 0.5, "Трат в этом месяце пока нет", ha="center", va="center", transform=ax2.transAxes)
    ax2.set_title("Куда уходят траты, ₽", fontsize=11, fontweight="bold", loc="left")

    fig.tight_layout()
    return _to_buf(fig)


def chart_waterfall(user_id: int) -> io.BytesIO:
    today, info, start, by_day, _ = collect_month(user_id)
    spent = sum(by_day.values())
    steps = [
        ("Обязательные", info["mandatory_total"]),
        ("Копилка", max(0.0, info["total_save_effective"])),
        ("Траты", spent),
    ]
    income = info["income"]
    final = income - sum(v for _, v in steps)

    fig = Figure(figsize=(8, 4.8))
    ax = fig.add_subplot(111)
    labels = ["Доход"] + [n for n, _ in steps] + ["Остаток"]
    ax.bar(0, income, color=BLUE, width=0.6)
    ax.text(0, income, fmt_num(income), ha="center", va="bottom", fontsize=10, fontweight="bold")

    level = income
    for i, (_, v) in enumerate(steps, start=1):
        new_level = level - v
        ax.bar(i, abs(level - new_level), bottom=min(level, new_level), color=ORANGE, width=0.6)
        ax.text(i, max(level, new_level), f"−{fmt_num(v)}", ha="center", va="bottom",
                fontsize=10, fontweight="bold")
        level = new_level

    last = len(labels) - 1
    ax.bar(last, abs(final), bottom=min(0, final), color=GREEN if final >= 0 else RED, width=0.6)
    ax.text(last, max(0, final), f"{'' if final >= 0 else '−'}{fmt_num(abs(final))}",
            ha="center", va="bottom", fontsize=10, fontweight="bold")
    ax.axhline(0, color=GRAY, linewidth=1)
    ax.set_xticks(range(len(labels)))
    ax.set_xticklabels(labels)
    _style(ax)
    top = max(income, 1)
    ax.set_ylim(min(0, final) - top * 0.05, top * 1.12)
    ax.set_title(
        f"{MONTHS_RU[today.month - 1]}: от дохода до остатка (на {today.strftime('%d.%m')}), ₽",
        fontsize=11, fontweight="bold", loc="left",
    )
    fig.tight_layout()
    return _to_buf(fig)


CHARTS = {
    "overview": (chart_overview, "📊 Траты по дням и банк сдачи"),
    "calendar": (chart_calendar, "🗓 Календарь месяца: зелёный день в лимите, красный перерасход"),
    "forecast": (chart_forecast, "🔮 Прогноз на конец месяца по текущему темпу трат"),
    "structure": (chart_structure, "🍩 Структура дохода и категории трат"),
    "waterfall": (chart_waterfall, "💧 От дохода до остатка"),
}


def build_chart(kind: str, user_id: int) -> io.BytesIO:
    return CHARTS[kind][0](user_id)


async def send_chart_kind(message, user_id: int, kind: str, edit: bool = False):
    try:
        buf = await asyncio.to_thread(build_chart, kind, user_id)
    except Exception:
        logger.exception("Ошибка построения графика %s", kind)
        await message.reply_text("⚠️ Не удалось построить график.", reply_markup=MAIN_KB)
        return

    caption = CHARTS[kind][1]
    if edit:
        try:
            await message.edit_media(
                media=InputMediaPhoto(media=buf, caption=caption),
                reply_markup=chart_kb(),
            )
            return
        except TelegramError as e:
            if "not modified" in str(e).lower():
                return
            buf.seek(0)
    await message.reply_photo(photo=buf, caption=caption, reply_markup=chart_kb())


# --- Фоновая задача: вечерний отчёт ---
async def daily_reminder_job(context: ContextTypes.DEFAULT_TYPE):
    now_dt = datetime.datetime.now(TZ)
    t_date = now_dt.date()
    now_min = now_dt.hour * 60 + now_dt.minute

    with db() as conn:
        user_rows = conn.execute(
            "SELECT DISTINCT user_id FROM settings WHERE key = 'daily_income'"
        ).fetchall()

    for r in user_rows:
        uid = r["user_id"]
        parsed = parse_time(get_setting(uid, "reminder_time", "21:00") or "21:00") or (21, 0)
        target_min = parsed[0] * 60 + parsed[1]

        # отчёт уходит после целевого времени, если в этот день ещё не отправлялся
        # (так он придёт и после простоя бота, а не потеряется)
        if now_min < target_min:
            continue
        if get_setting(uid, "last_report") == t_date.isoformat():
            continue
        if get_setting(uid, "blocked") == "1":
            continue

        sent = False
        try:
            text, markup = build_day_report(uid, t_date)
            header = "⏰ *Ваш автоматический вечерний отчёт за день!*\n\n"
            await context.bot.send_message(
                chat_id=uid,
                text=header + text,
                parse_mode="Markdown",
                reply_markup=markup or MAIN_KB,
            )
            sent = True
            set_setting(uid, "last_report", t_date.isoformat())  # только после успешной отправки
            logger.info("Вечерний отчёт отправлен пользователю %s", uid)
        except Forbidden:
            logger.info("Пользователь %s заблокировал бота — отчёты приостановлены", uid)
            set_setting(uid, "last_report", t_date.isoformat())
            set_setting(uid, "blocked", "1")
        except Exception as e:
            fails = context.application.bot_data.setdefault("report_fails", {})
            day, cnt = fails.get(uid, (None, 0))
            cnt = cnt + 1 if day == t_date.isoformat() else 1
            fails[uid] = (t_date.isoformat(), cnt)
            logger.error("Не удалось отправить отчёт пользователю %s (попытка %s): %s", uid, cnt, e)
            if cnt >= 5:  # не долбим бесконечно
                set_setting(uid, "last_report", t_date.isoformat())

        if sent and get_setting(uid, "report_chart", "1") == "1":
            try:
                buf = await asyncio.to_thread(build_chart, "overview", uid)
                await context.bot.send_photo(
                    chat_id=uid, photo=buf, caption=CHARTS["overview"][1], reply_markup=chart_kb()
                )
            except Exception as e:
                logger.error("Не удалось отправить график пользователю %s: %s", uid, e)
        await asyncio.sleep(0.05)


# --- Команды ---
async def start_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    ensure_start_day(user_id)
    if get_setting(user_id, "blocked") == "1":
        set_setting(user_id, "blocked", "0")
    context.user_data.clear()
    context.user_data["wizard_step"] = "wizard_mode"
    welcome = (
        "👋 *Привет! Добро пожаловать в твой персональный бот для бюджетирования.*\n\n"
        "🧙‍♂️ Давай настроим твой личный бюджет за пару шагов.\n\n"
        "📥 *Шаг 1 из 5:* Как ты получаешь деньги?"
    )
    await update.message.reply_text(welcome, parse_mode="Markdown", reply_markup=mode_markup("wiz_mode_"))


async def cancel_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    await update.message.reply_text("Ввод отменён.", reply_markup=MAIN_KB)


async def show_today(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    text, markup = build_day_report(user_id, today_date())
    await update.message.reply_text(text, parse_mode="Markdown", reply_markup=markup or MAIN_KB)


async def show_settings(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    context.user_data.clear()
    await update.message.reply_text(
        build_settings_text(user_id),
        parse_mode="Markdown",
        reply_markup=settings_inline_kb(user_id),
    )


async def show_week(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    end_d = today_date()
    start_d = end_d - datetime.timedelta(days=6)
    info = get_budget_info(user_id, end_d)
    lims = limits_between(user_id, start_d, end_d)
    with db() as conn:
        rows = conn.execute(
            "SELECT spent_date, SUM(amount) AS total FROM expenses "
            "WHERE user_id = ? AND spent_date >= ? AND spent_date <= ? GROUP BY spent_date",
            (user_id, start_d.isoformat(), end_d.isoformat()),
        ).fetchall()
    by_day = {r["spent_date"]: r["total"] for r in rows}

    lines = ["📆 *Траты за последние 7 дней:*\n"]
    total_week = 0.0
    for i in range(7):
        cur = start_d + datetime.timedelta(days=i)
        sp = by_day.get(cur.isoformat(), 0.0)
        total_week += sp
        diff = lims.get(cur, info["daily_limit"]) - sp
        icon = "🟢" if diff >= 0 else "🔴"
        lines.append(f"{icon} *{cur.strftime('%d.%m')}*: {fmt_money(sp)} (остаток: {fmt_money(diff)})")

    lines.append(f"\n💸 *Всего за 7 дней:* {fmt_money(total_week)}")
    await update.message.reply_text("\n".join(lines), parse_mode="Markdown", reply_markup=MAIN_KB)


async def show_month(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    now_d = today_date()
    month_prefix = f"{now_d.year:04d}-{now_d.month:02d}-"
    info = get_budget_info(user_id, now_d)

    with db() as conn:
        day_rows = conn.execute(
            "SELECT spent_date, SUM(amount) AS total FROM expenses "
            "WHERE user_id = ? AND spent_date LIKE ? GROUP BY spent_date ORDER BY spent_date",
            (user_id, month_prefix + "%"),
        ).fetchall()
        top_rows = conn.execute(
            "SELECT category, SUM(amount) AS total FROM expenses "
            "WHERE user_id = ? AND spent_date LIKE ? GROUP BY category ORDER BY total DESC LIMIT 5",
            (user_id, month_prefix + "%"),
        ).fetchall()

    expenses_spent = sum(r["total"] for r in day_rows)
    total_spent = expenses_spent + info["mandatory_total"]
    change_pool = calc_change_pool(user_id, now_d)

    lines = [f"📅 *Сводка за {now_d.strftime('%m.%Y')}*\n"]

    if info["mandatory_list"]:
        lines.append("🏠 *Обязательные расходы (учтены):*")
        for m in info["mandatory_list"]:
            lines.append(f"• {esc(m['name'])}: {fmt_money(m['amount'])}")
        lines.append("")

    if day_rows:
        lines.append("🛍 *Обычные траты по дням:*")
        for r in day_rows[-10:]:
            d_fmt = datetime.date.fromisoformat(r["spent_date"]).strftime("%d.%m")
            lines.append(f"• {d_fmt}: *{fmt_money(r['total'])}*")
    else:
        lines.append("_В этом месяце обычных трат пока нет_")

    lines.append(f"\n💸 *Всего потрачено (с учётом обязательных):* *{fmt_money(total_spent)}*")
    lines.append(f"🪙 *Банк сдачи:* {fmt_money(change_pool)}")
    lines.append(f"💰 *Копилка (всего):* {fmt_money(info['piggy_total'])}")

    if top_rows:
        lines.append("\n🏆 *Топ-5 категорий расходов:*")
        for i, r in enumerate(top_rows, 1):
            lines.append(f"{i}. {esc(r['category'])} — {fmt_money(r['total'])}")

    await update.message.reply_text("\n".join(lines), parse_mode="Markdown", reply_markup=MAIN_KB)


async def send_chart(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await send_chart_kind(update.message, update.effective_user.id, "overview")


async def send_excel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    try:
        buf = await asyncio.to_thread(update_excel_file, user_id)
        await update.message.reply_document(
            document=buf,
            filename=f"Бюджет_{today_date().strftime('%m_%Y')}.xlsx",
            caption="📥 Ваша персональная таблица бюджета с деталями трат",
            reply_markup=MAIN_KB,
        )
    except Exception:
        logger.exception("Ошибка при формировании Excel:")
        await update.message.reply_text("⚠️ Не удалось сформировать Excel. Попробуйте позже.", reply_markup=MAIN_KB)


async def undo_last(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    with db() as conn:
        exp = conn.execute(
            "SELECT id, amount, category, spent_date, created_at FROM expenses "
            "WHERE user_id = ? ORDER BY id DESC LIMIT 1",
            (user_id,),
        ).fetchone()
        top = conn.execute(
            "SELECT id, amount, comment, topup_date, created_at FROM savings_topups "
            "WHERE user_id = ? ORDER BY id DESC LIMIT 1",
            (user_id,),
        ).fetchone()
    if not exp and not top:
        await update.message.reply_text("Операций для отмены пока нет.", reply_markup=MAIN_KB)
        return

    if top and (not exp or (top["created_at"] or "") >= (exp["created_at"] or "")):
        freeze_past_limits(user_id)
        with db() as conn:
            conn.execute("DELETE FROM savings_topups WHERE id = ? AND user_id = ?", (top["id"], user_id))
        sign = "+" if top["amount"] >= 0 else "−"
        await update.message.reply_text(
            f"↩️ Отменена операция с копилкой: *{sign}{fmt_money(abs(top['amount']))}* "
            f"({esc(top['comment'])}, {top['topup_date']})",
            parse_mode="Markdown",
            reply_markup=MAIN_KB,
        )
        return

    delete_expense(user_id, exp["id"])
    await update.message.reply_text(
        f"↩️ Удалена последняя трата: *{fmt_money(exp['amount'])}* ({esc(exp['category'])}, {exp['spent_date']})",
        parse_mode="Markdown",
        reply_markup=MAIN_KB,
    )


def build_history(user_id: int):
    with db() as conn:
        rows = conn.execute(
            "SELECT id, spent_date, amount, category FROM expenses WHERE user_id = ? "
            "ORDER BY spent_date DESC, id DESC LIMIT 12",
            (user_id,),
        ).fetchall()
    if not rows:
        return "Трат пока нет. Запишите первую, например: `450 продукты`.", None
    lines = ["🧾 *Последние траты* (нажмите, чтобы удалить):", ""]
    buttons = []
    for r in rows:
        d = datetime.date.fromisoformat(r["spent_date"]).strftime("%d.%m")
        lines.append(f"• {d} — *{fmt_money(r['amount'])}* {esc(r['category'])}")
        title = f"❌ {d} {int(r['amount'])}₽ {r['category']}"
        buttons.append([InlineKeyboardButton(title[:40], callback_data=f"hdelexp_{r['id']}")])
    return "\n".join(lines), InlineKeyboardMarkup(buttons)


async def show_history(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text, markup = build_history(update.effective_user.id)
    await update.message.reply_text(text, parse_mode="Markdown", reply_markup=markup or MAIN_KB)


def piggy_undo_kb(topup_id: int):
    return InlineKeyboardMarkup([[InlineKeyboardButton("↩️ Отменить", callback_data=f"deltop_{topup_id}")]])


def build_piggy_history(user_id: int):
    info = get_budget_info(user_id)
    with db() as conn:
        rows = conn.execute(
            "SELECT id, topup_date, amount, comment FROM savings_topups "
            "WHERE user_id = ? ORDER BY id DESC LIMIT 10",
            (user_id,),
        ).fetchall()
    lines = [
        f"🏦 *Копилка (всего):* {fmt_money(info['piggy_total'])}",
        f"_Плановые отчисления: {fmt_money(info['save'])} в месяц_",
        "",
    ]
    buttons = []
    if rows:
        lines.append("*Последние операции* (нажмите, чтобы отменить):")
        for r in rows:
            d = datetime.date.fromisoformat(r["topup_date"]).strftime("%d.%m")
            sign = "+" if r["amount"] >= 0 else "−"
            lines.append(f"• {d} {sign}{fmt_money(abs(r['amount']))} — {esc(r['comment'])}")
            title = f"❌ {d} {sign}{int(abs(r['amount']))}₽ {r['comment']}"
            buttons.append([InlineKeyboardButton(title[:40], callback_data=f"hdeltop_{r['id']}")])
    else:
        lines.append("_Операций пока нет. Пополнить: кнопка «🏦 В копилку» или `+500`._")
    return "\n".join(lines), InlineKeyboardMarkup(buttons) if buttons else None


async def show_piggy(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text, markup = build_piggy_history(update.effective_user.id)
    await update.message.reply_text(text, parse_mode="Markdown", reply_markup=markup or MAIN_KB)


async def piggy_add(update: Update, context: ContextTypes.DEFAULT_TYPE, user_id: int, amount: float, comment: str):
    freeze_past_limits(user_id)
    with db() as conn:
        tid = conn.execute(
            "INSERT INTO savings_topups (user_id, topup_date, amount, comment) VALUES (?, ?, ?, ?)",
            (user_id, today_date().isoformat(), amount, comment),
        ).lastrowid
    context.user_data.clear()
    info = get_budget_info(user_id)
    await update.message.reply_text(
        f"💰 В копилку зачислено: *+{fmt_money(amount)}* ({esc(comment)})!\n"
        f"🏦 Всего в копилке: *{fmt_money(info['piggy_total'])}*\n"
        f"_Лимит на сегодня и дальше: {fmt_money(info['daily_limit'])} (прошлые дни не менялись)._",
        parse_mode="Markdown",
        reply_markup=piggy_undo_kb(tid),
    )


async def help_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "ℹ️ *Как пользоваться*\n\n"
        "Трата: `450 продукты`, `такси 300 вчера`, `1 500 одежда 28.10`\n"
        "Пополнить копилку: `+500 подработка`\n"
        "Снять из копилки — кнопка «💸 Из копилки» (на дневной лимит не влияет)\n\n"
        "Команды: /today /week /month /history /chart /excel /piggy /settings /undo /cancel",
        parse_mode="Markdown",
        reply_markup=MAIN_KB,
    )


async def post_init(app):
    try:
        await app.bot.set_my_commands([
            BotCommand("today", "Траты за сегодня"),
            BotCommand("week", "Траты за 7 дней"),
            BotCommand("month", "Сводка за месяц"),
            BotCommand("history", "Последние траты"),
            BotCommand("chart", "График"),
            BotCommand("excel", "Скачать Excel"),
            BotCommand("piggy", "Копилка и её история"),
            BotCommand("settings", "Настройки"),
            BotCommand("undo", "Отменить последнюю операцию"),
            BotCommand("help", "Как вводить траты"),
        ])
    except TelegramError:
        logger.warning("Не удалось установить меню команд")


async def error_handler(update: object, context: ContextTypes.DEFAULT_TYPE):
    logger.error("Необработанная ошибка", exc_info=context.error)
    if isinstance(update, Update) and update.effective_message:
        try:
            await update.effective_message.reply_text("⚠️ Что-то пошло не так. Попробуйте ещё раз.")
        except TelegramError:
            pass


# --- Inline-кнопки ---
async def refresh_settings(query, user_id: int, prefix: str = ""):
    text = prefix + build_settings_text(user_id)
    try:
        await query.message.edit_text(text, parse_mode="Markdown", reply_markup=settings_inline_kb(user_id))
    except TelegramError:
        await query.message.reply_text(text, parse_mode="Markdown", reply_markup=settings_inline_kb(user_id))


async def callback_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    data = query.data
    user_id = update.effective_user.id

    if data.startswith("chart_"):
        kind = data.removeprefix("chart_")
        if kind in CHARTS:
            await send_chart_kind(query.message, user_id, kind, edit=True)
        return

    if data == "restart_bot_action":
        context.user_data.clear()
        await refresh_settings(query, user_id, "🧹 *Текущий ввод сброшен.*\n\n")
        return

    if data == "full_reset_action":
        markup = InlineKeyboardMarkup([[
            InlineKeyboardButton("🗑 Да, удалить всё", callback_data="full_reset_confirm"),
            InlineKeyboardButton("Отмена", callback_data="full_reset_cancel"),
        ]])
        await query.message.reply_text(
            "⚠️ *Удалить все настройки, траты и копилку?* Это нельзя отменить.",
            parse_mode="Markdown",
            reply_markup=markup,
        )
        return

    if data == "full_reset_cancel":
        await query.message.edit_text("Отменено, данные не тронуты.")
        return

    if data == "full_reset_confirm":
        with db() as conn:
            conn.execute("DELETE FROM settings WHERE user_id = ?", (user_id,))
            conn.execute("DELETE FROM mandatory WHERE user_id = ?", (user_id,))
            conn.execute("DELETE FROM expenses WHERE user_id = ?", (user_id,))
            conn.execute("DELETE FROM savings_topups WHERE user_id = ?", (user_id,))
            conn.execute("DELETE FROM day_limits WHERE user_id = ?", (user_id,))
        context.user_data.clear()
        context.user_data["wizard_step"] = "wizard_mode"
        await query.message.edit_text(
            "⚠️ *Все ваши данные, настройки, траты и копилка удалены.*\n\n"
            "🧙‍♂️ Давайте настроим бюджет заново с чистого листа.\n\n"
            "📥 *Шаг 1 из 5:* Как вы получаете деньги?",
            parse_mode="Markdown",
            reply_markup=mode_markup("wiz_mode_"),
        )
        return

    if data.startswith("wiz_mode_"):
        chosen = data.removeprefix("wiz_mode_")
        if chosen not in INCOME_MODES:
            return
        set_setting(user_id, "income_mode", chosen)
        context.user_data["wizard_step"] = "wizard_income"
        await query.message.edit_text(
            "🏦 *Шаг 2 из 5:* Введи сумму дохода (число следующим сообщением, например: `1000` или `50000`):",
            parse_mode="Markdown",
        )
        return

    if data == "set_daily_income":
        context.user_data["state"] = "wait_daily_income"
        mode = get_setting(user_id, "income_mode", "daily")
        unit = INCOME_UNIT_PROMPT.get(mode, "в день")
        await query.message.reply_text(
            f"🏦 Введите сумму вашего дохода *{unit}* (например: `1000` или `50050`):",
            parse_mode="Markdown",
        )

    elif data == "set_income_mode":
        await query.message.reply_text("🔄 Выберите график получения дохода:", reply_markup=mode_markup("mode_"))

    elif data.startswith("mode_"):
        m_type = data.removeprefix("mode_")
        if m_type not in INCOME_MODES:
            return
        set_setting(user_id, "income_mode", m_type)
        await refresh_settings(query, user_id, "✅ График дохода изменён!\n\n")

    elif data == "toggle_report_chart":
        cur = get_setting(user_id, "report_chart", "1")
        set_setting(user_id, "report_chart", "0" if cur == "1" else "1")
        await refresh_settings(query, user_id)

    elif data == "set_days_count":
        context.user_data["state"] = "wait_days_count"
        await query.message.reply_text(
            "📅 Сколько дней учитывать в месяце?\n"
            "• Введите число (например: `30`, `31`)\n"
            "• Или отправьте слово `авто`, чтобы бот брал реальное количество дней из календаря.",
            parse_mode="Markdown",
        )

    elif data == "set_reminder_time":
        context.user_data["state"] = "wait_reminder_time"
        await query.message.reply_text(
            f"⏰ Введите время ({TZ_LABEL}) для вечернего отчёта (формат `ЧЧ:ММ`, например: `21:30`):",
            parse_mode="Markdown",
        )

    elif data == "set_save":
        context.user_data["state"] = "wait_save"
        await query.message.reply_text(
            "💰 Введите сумму плановой копилки за месяц (например: `5000`, или `0`, чтобы отключить):",
            parse_mode="Markdown",
        )

    elif data == "add_mand":
        context.user_data["state"] = "wait_mand"
        await query.message.reply_text(
            "➕ Напишите название и сумму обязательного расхода через пробел (например: `Аренда 15000`):",
            parse_mode="Markdown",
        )

    elif data == "del_mand_menu":
        info = get_budget_info(user_id)
        if not info["mandatory_list"]:
            await query.message.reply_text("Ваш список обязательных расходов пуст.")
            return
        btns = [
            [InlineKeyboardButton(f"❌ {r['name']} ({fmt_money(r['amount'])})", callback_data=f"delmand_{r['id']}")]
            for r in info["mandatory_list"]
        ]
        await query.message.reply_text(
            "🗑 Нажмите на обязательный расход, который нужно удалить:",
            reply_markup=InlineKeyboardMarkup(btns),
        )

    elif data.startswith("delmand_"):
        try:
            mand_id = int(data.split("_")[1])
        except (ValueError, IndexError):
            return
        freeze_past_limits(user_id)
        with db() as conn:
            conn.execute("DELETE FROM mandatory WHERE id = ? AND user_id = ?", (mand_id, user_id))
        await refresh_settings(query, user_id, "✅ Обязательный расход удалён!\n\n")

    elif data == "restart_wizard":
        context.user_data.clear()
        context.user_data["wizard_step"] = "wizard_mode"
        await query.message.edit_text(
            "🔄 *Запуск мастера настройки*\n\n📥 *Шаг 1 из 5:* Как ты получаешь деньги?",
            parse_mode="Markdown",
            reply_markup=mode_markup("wiz_mode_"),
        )

    elif data.startswith("hdelexp_"):
        try:
            exp_id = int(data.split("_")[1])
        except (ValueError, IndexError):
            return
        delete_expense(user_id, exp_id)
        text, markup = build_history(user_id)
        try:
            await query.message.edit_text(text, parse_mode="Markdown", reply_markup=markup)
        except TelegramError:
            pass

    elif data == "piggy_history":
        text, markup = build_piggy_history(user_id)
        await query.message.reply_text(text, parse_mode="Markdown", reply_markup=markup)

    elif data.startswith("deltop_") or data.startswith("hdeltop_"):
        try:
            tid = int(data.split("_")[1])
        except (ValueError, IndexError):
            return
        freeze_past_limits(user_id)
        with db() as conn:
            deleted = conn.execute(
                "DELETE FROM savings_topups WHERE id = ? AND user_id = ?", (tid, user_id)
            ).rowcount
        if data.startswith("hdeltop_"):
            text, markup = build_piggy_history(user_id)
            try:
                await query.message.edit_text(text, parse_mode="Markdown", reply_markup=markup)
            except TelegramError:
                pass
        else:
            msg = "↩️ Операция с копилкой отменена." if deleted else "ℹ️ Эта операция уже отменена."
            try:
                await query.message.edit_text(msg)
            except TelegramError:
                pass

    elif data.startswith("delexp_"):
        try:
            exp_id = int(data.split("_")[1])
        except (ValueError, IndexError):
            return
        spent_date_str = delete_expense(user_id, exp_id)
        if not spent_date_str:
            await query.message.reply_text("ℹ️ Эта трата уже удалена.")
            return
        target_d = datetime.date.fromisoformat(spent_date_str)
        text, markup = build_day_report(user_id, target_d)
        try:
            await query.message.edit_text(text, parse_mode="Markdown", reply_markup=markup)
        except TelegramError:
            pass


# --- Текстовые сообщения ---
async def text_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = (update.message.text or "").strip()
    user_id = update.effective_user.id
    if get_setting(user_id, "blocked") == "1":
        set_setting(user_id, "blocked", "0")

    # нажатие кнопки меню всегда прерывает незавершённый ввод
    if text in MENU_TEXTS:
        context.user_data.clear()

    wizard_step = context.user_data.get("wizard_step")
    if wizard_step:
        if wizard_step == "wizard_income":
            val = parse_amount(text)
            if val is None:
                await update.message.reply_text("⚠️ Введи положительное число, например: `1000`", parse_mode="Markdown")
                return
            set_setting(user_id, "daily_income", val)
            context.user_data["wizard_step"] = "wizard_save"
            await update.message.reply_text(
                "💰 *Шаг 3 из 5:* Сколько откладывать в плановую копилку за месяц? (например: `5000` или `0`):",
                parse_mode="Markdown",
            )
            return

        if wizard_step == "wizard_save":
            val = parse_amount(text, allow_zero=True)
            if val is None:
                await update.message.reply_text("⚠️ Введи число, например: `5000`", parse_mode="Markdown")
                return
            set_setting(user_id, "save", val)
            context.user_data["wizard_step"] = "wizard_mand"
            await update.message.reply_text(
                "🏠 *Шаг 4 из 5:* Добавь первый *обязательный расход* (название и сумма, например: "
                "`Аренда 15000`, или напиши `нет`, чтобы пропустить):",
                parse_mode="Markdown",
            )
            return

        if wizard_step == "wizard_mand":
            if text.lower() not in ["нет", "0", "пропустить", "-"]:
                parsed = parse_label_amount(text)
                if not parsed:
                    await update.message.reply_text(
                        "⚠️ Напиши название и сумму, например: `Аренда 15000`, или `нет`.",
                        parse_mode="Markdown",
                    )
                    return
                amount, name = parsed
                upsert_mandatory(user_id, name, amount)
            context.user_data["wizard_step"] = "wizard_time"
            await update.message.reply_text(
                f"⏰ *Шаг 5 из 5:* В какое *время ({TZ_LABEL})* присылать вечерний отчёт? (например: `21:00`):",
                parse_mode="Markdown",
            )
            return

        if wizard_step == "wizard_time":
            t = parse_time(text)
            if not t:
                await update.message.reply_text("⚠️ Введи время в формате `ЧЧ:ММ`, например: `21:00`", parse_mode="Markdown")
                return
            save_reminder_time(user_id, *t)
            context.user_data.clear()
            await update.message.reply_text(
                "🎉 *Настройка завершена!*\n\n"
                "Чтобы записать трату, напиши, например: `450 продукты`.",
                parse_mode="Markdown",
                reply_markup=MAIN_KB,
            )
            return

    if text == "📊 Сегодня":
        await show_today(update, context)
        return
    if text == "📅 Месяц":
        await show_month(update, context)
        return
    if text == "🏦 В копилку":
        context.user_data["state"] = "wait_piggy"
        await update.message.reply_text(
            "💰 *Пополнение плановой копилки*\n\n"
            "Сколько хочешь отложить? (например: `500 Подработка` или просто `500`):",
            parse_mode="Markdown",
        )
        return
    if text == "💸 Из копилки":
        context.user_data["state"] = "wait_withdraw_piggy"
        await update.message.reply_text(
            "💸 *Снятие из копилки*\n\nВведите сумму и причину, например: `3000 ремонт телефона`:",
            parse_mode="Markdown",
        )
        return
    if text == "📈 График":
        await send_chart(update, context)
        return
    if text == "📥 Скачать Excel":
        await send_excel(update, context)
        return
    if text == "⚙️ Настройки":
        await show_settings(update, context)
        return
    if text == "↩️ Отменить последнюю":
        await undo_last(update, context)
        return

    state = context.user_data.get("state")

    if state == "wait_piggy":
        parsed = parse_amount_comment(text, "Пополнение")
        if not parsed:
            await update.message.reply_text("⚠️ Введи сумму числом, например: `500 Подработка`", parse_mode="Markdown")
            return
        await piggy_add(update, context, user_id, *parsed)
        return

    if state == "wait_withdraw_piggy":
        parsed = parse_amount_comment(text, "Снятие")
        if not parsed:
            await update.message.reply_text("⚠️ Введи сумму и причину, например: `3000 ремонт телефона`", parse_mode="Markdown")
            return
        amount, comment = parsed
        info = get_budget_info(user_id)
        if amount > info["piggy_total"] + 0.005:
            await update.message.reply_text(
                f"⚠️ В копилке только *{fmt_money(max(0.0, info['piggy_total']))}*, а вы хотите снять "
                f"*{fmt_money(amount)}*. Введите сумму поменьше или нажмите /cancel.",
                parse_mode="Markdown",
            )
            return
        with db() as conn:
            tid = conn.execute(
                "INSERT INTO savings_topups (user_id, topup_date, amount, comment) VALUES (?, ?, ?, ?)",
                (user_id, today_date().isoformat(), -amount, f"Снятие: {comment}"),
            ).lastrowid
        context.user_data.clear()
        info = get_budget_info(user_id)
        await update.message.reply_text(
            f"💸 Снято из копилки: *-{fmt_money(amount)}* ({esc(comment)})!\n"
            f"🏦 Осталось в копилке: *{fmt_money(info['piggy_total'])}*\n"
            "_На дневной лимит и банк сдачи снятие не влияет._",
            parse_mode="Markdown",
            reply_markup=piggy_undo_kb(tid),
        )
        return

    if state == "wait_daily_income":
        val = parse_amount(text)
        if val is None:
            await update.message.reply_text("⚠️ Введите положительное число, например: `1000`", parse_mode="Markdown")
            return
        set_setting(user_id, "daily_income", val)
        context.user_data.clear()
        await update.message.reply_text(
            "✅ Доход обновлён!\n\n" + build_settings_text(user_id),
            parse_mode="Markdown",
            reply_markup=settings_inline_kb(user_id),
        )
        return

    if state == "wait_days_count":
        if text.lower() == "авто":
            freeze_past_limits(user_id)
            with db() as conn:
                conn.execute(
                    "DELETE FROM settings WHERE user_id = ? AND key = 'custom_days_in_month'", (user_id,)
                )
            context.user_data.clear()
            await update.message.reply_text(
                "✅ Дней в месяце теперь берётся *автоматически по календарю*.\n\n" + build_settings_text(user_id),
                parse_mode="Markdown",
                reply_markup=settings_inline_kb(user_id),
            )
            return
        try:
            days_val = int(text)
        except ValueError:
            await update.message.reply_text("⚠️ Введи число или слово `авто`.", parse_mode="Markdown")
            return
        if not 1 <= days_val <= 31:
            await update.message.reply_text("⚠️ Введи число от 1 до 31 или слово `авто`.", parse_mode="Markdown")
            return
        set_setting(user_id, "custom_days_in_month", days_val)
        context.user_data.clear()
        await update.message.reply_text(
            f"✅ Фиксированное количество дней в месяце: *{days_val}*\n\n" + build_settings_text(user_id),
            parse_mode="Markdown",
            reply_markup=settings_inline_kb(user_id),
        )
        return

    if state == "wait_reminder_time":
        t = parse_time(text)
        if not t:
            await update.message.reply_text(
                "⚠️ Неверный формат. Введи время как `ЧЧ:ММ` (часы 0-23, минуты 0-59), например: `21:30`",
                parse_mode="Markdown",
            )
            return
        save_reminder_time(user_id, *t)
        context.user_data.clear()
        await update.message.reply_text(
            f"⏰ Время вечернего отчёта изменено на *{t[0]:02d}:{t[1]:02d} ({TZ_LABEL})*!\n\n" + build_settings_text(user_id),
            parse_mode="Markdown",
            reply_markup=settings_inline_kb(user_id),
        )
        return

    if state == "wait_save":
        val = parse_amount(text, allow_zero=True)
        if val is None:
            await update.message.reply_text("⚠️ Введи число, например: `5000`", parse_mode="Markdown")
            return
        set_setting(user_id, "save", val)
        context.user_data.clear()
        await update.message.reply_text(
            f"💰 Плановая копилка обновлена: *{fmt_money(val)}*\n\n" + build_settings_text(user_id),
            parse_mode="Markdown",
            reply_markup=settings_inline_kb(user_id),
        )
        return

    if state == "wait_mand":
        parsed = parse_label_amount(text)
        if not parsed:
            await update.message.reply_text("⚠️ Напиши название и сумму, например: `Аренда 15000`", parse_mode="Markdown")
            return
        amount, name = parsed
        updated = upsert_mandatory(user_id, name, amount)
        context.user_data.clear()
        await update.message.reply_text(
            f"✅ {'Обновлён' if updated else 'Добавлен'} обязательный расход: *{esc(name)} — {fmt_money(amount)}*\n\n" + build_settings_text(user_id),
            parse_mode="Markdown",
            reply_markup=settings_inline_kb(user_id),
        )
        return

    if text.startswith("+"):
        parsed = parse_amount_comment(text[1:], "Пополнение")
        if parsed:
            await piggy_add(update, context, user_id, *parsed)
            return

    parsed = parse_expense_message(text)
    if not parsed:
        await update.message.reply_text(
            "🤔 Не понял сообщение.\nЧтобы записать трату: `450 продукты`.",
            parse_mode="Markdown",
            reply_markup=MAIN_KB,
        )
        return

    amount, category, spent_d = parsed
    if spent_d > today_date():
        await update.message.reply_text(
            f"⚠️ Дата {spent_d.strftime('%d.%m.%Y')} ещё не наступила, трату не записал. "
            "Укажите сегодня, вчера или прошедшую дату.",
            reply_markup=MAIN_KB,
        )
        return
    ensure_start_day(user_id)
    with db() as conn:
        conn.execute(
            "INSERT INTO expenses (user_id, spent_date, amount, category) VALUES (?, ?, ?, ?)",
            (user_id, spent_d.isoformat(), amount, category),
        )

    report_text, markup = build_day_report(user_id, spent_d)
    await update.message.reply_text(
        f"✅ Записано: *{fmt_money(amount)}* — {esc(category)} ({spent_d.strftime('%d.%m')})\n\n{report_text}",
        parse_mode="Markdown",
        reply_markup=markup or MAIN_KB,
    )


def main():
    if not BOT_TOKEN:
        raise RuntimeError("Не указан BOT_TOKEN в файле .env")
    init_db()
    logger.info("Бот запущен. Часовой пояс: %s, база: %s", TZ_NAME, DB_PATH)

    persistence = PicklePersistence(filepath=STATE_PATH)
    app = Application.builder().token(BOT_TOKEN).persistence(persistence).post_init(post_init).build()

    if app.job_queue:
        app.job_queue.run_repeating(daily_reminder_job, interval=60, first=10)

    app.add_handler(CommandHandler("start", start_cmd))
    app.add_handler(CommandHandler("cancel", cancel_cmd))
    app.add_handler(CommandHandler("today", show_today))
    app.add_handler(CommandHandler("week", show_week))
    app.add_handler(CommandHandler("month", show_month))
    app.add_handler(CommandHandler("chart", send_chart))
    app.add_handler(CommandHandler("excel", send_excel))
    app.add_handler(CommandHandler("settings", show_settings))
    app.add_handler(CommandHandler("undo", undo_last))
    app.add_handler(CommandHandler("piggy", show_piggy))
    app.add_handler(CommandHandler("history", show_history))
    app.add_handler(CommandHandler("help", help_cmd))

    app.add_handler(CallbackQueryHandler(callback_handler))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, text_handler))
    app.add_error_handler(error_handler)

    app.run_polling()


if __name__ == "__main__":
    main()
